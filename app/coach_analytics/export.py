"""
Native .xlsx export for Coach Analytics (feedback build item E — billing).

Workbooks are built server-side with openpyxl and returned base64-encoded inside
the standard {data} envelope, so the locked frontend api client can carry them
without a binary code path (the frontend decodes to a Blob and downloads).

  • per-coach workbook  — Summary, Sessions held, 1-on-1 log (billable),
                          Player performance
  • all-coaches workbook — one summary row per coach

Aggregation is NOT duplicated here — both builders read the same controllers
(list_coach_analytics / get_coach_analytics) the API already serves.
"""

import base64
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.coach_analytics.controllers import (
    list_coach_analytics, get_coach_analytics,
)

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="0B2A4A")
_TITLE_FONT = Font(bold=True, size=14)


def _slug(text):
    return re.sub(r"[^A-Za-z0-9]+", "-", (text or "").strip()).strip("-").lower() or "coach"


def _pct(rate):
    return "" if rate is None else f"{round(rate * 100)}%"


def _hcp(change):
    # Positive == improvement (index dropped). Mirror the UI's sign sense.
    if change is None:
        return ""
    if change > 0:
        return f"-{abs(change):.1f} (improved)"
    if change < 0:
        return f"+{abs(change):.1f}"
    return "0.0"


def _write_header(ws, row, headers):
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="left")


def _autosize(ws, max_width=48):
    for col_cells in ws.columns:
        length = max(
            (len(str(c.value)) for c in col_cells if c.value is not None),
            default=0,
        )
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), max_width)


def _to_b64(wb):
    buf = BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


# ── per-coach workbook ────────────────────────────────────────────────────────

def build_coach_workbook(coach_id, date_from=None, date_to=None):
    """Returns ({filename, mime, content_base64}, None) or (None, error)."""
    detail, err = get_coach_analytics(coach_id, date_from=date_from, date_to=date_to)
    if err:
        return None, err

    window = detail["window"]
    wb = Workbook()

    # Sheet 1 — Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Coach analytics — {detail['coach_name']}"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Period: {window['date_from']} to {window['date_to']}"
    rows = [
        ("Assigned juniors", detail["junior_count"]),
        ("Sessions run", detail["sessions_run"]),
        ("Billable 1-on-1 sessions", detail["billable_one_on_one"]["count"]),
        ("Billable 1-on-1 attendees", detail["billable_one_on_one"]["attendees"]),
        ("Attendance present", detail["attendance"]["present"]),
        ("Attendance total", detail["attendance"]["total"]),
        ("Attendance rate", _pct(detail["attendance"]["rate"])),
        ("Average level", detail["avg_current_level"] if detail["avg_current_level"] is not None else ""),
        ("Average handicap change", _hcp(detail["avg_handicap_change"])),
        ("Juniors with handicap", detail["juniors_with_handicap"]),
        ("Evaluations signed", detail["evaluations"]["total"]),
        ("  — below expectation", detail["evaluations"]["below_expectation"]),
        ("  — meeting expectation", detail["evaluations"]["meeting_expectation"]),
        ("  — exceeding expectation", detail["evaluations"]["exceeding_expectation"]),
        ("  — recommended for promotion", detail["evaluations"]["move_next_level"]),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    _autosize(ws)

    # Sheet 2 — Sessions held
    ws = wb.create_sheet("Sessions held")
    _write_header(ws, 1, ["Date", "Type", "Focus", "Present", "Total", "Billable"])
    for r, s in enumerate(detail["sessions"], start=2):
        ws.cell(row=r, column=1, value=s["date"])
        ws.cell(row=r, column=2, value=s["session_type"])
        ws.cell(row=r, column=3, value=s["title"] or "")
        ws.cell(row=r, column=4, value=s["present"])
        ws.cell(row=r, column=5, value=s["total"])
        ws.cell(row=r, column=6, value="Yes" if s["billable"] else "")
    _autosize(ws)

    # Sheet 3 — 1-on-1 log (billable)
    ws = wb.create_sheet("1-on-1 log (billable)")
    _write_header(ws, 1, ["Date", "Focus", "Attendees"])
    billable = [s for s in detail["sessions"] if s["billable"]]
    for r, s in enumerate(billable, start=2):
        ws.cell(row=r, column=1, value=s["date"])
        ws.cell(row=r, column=2, value=s["title"] or "")
        ws.cell(row=r, column=3, value=s["present"])
    total_row = len(billable) + 2
    ws.cell(row=total_row, column=2, value="Total billable sessions").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=len(billable)).font = Font(bold=True)
    _autosize(ws)

    # Sheet 4 — Player performance / growth
    ws = wb.create_sheet("Player performance")
    _write_header(ws, 1, [
        "Junior", "Band", "Level", "Handicap", "Handicap change",
        "Latest assessment", "Recommendation", "Evaluation month",
    ])
    for r, j in enumerate(detail["juniors"], start=2):
        ws.cell(row=r, column=1, value=j["full_name"])
        ws.cell(row=r, column=2, value=j["band_label"] or "")
        ws.cell(row=r, column=3, value=j["current_level"] if j["current_level"] is not None else "")
        ws.cell(row=r, column=4, value=j["handicap_index"] if j["has_handicap"] and j["handicap_index"] is not None else "")
        ws.cell(row=r, column=5, value=_hcp(j["handicap_change"]))
        ws.cell(row=r, column=6, value=j["latest_assessment"] or "")
        ws.cell(row=r, column=7, value=j["latest_recommendation"] or "")
        ws.cell(row=r, column=8, value=j["latest_eval_month"] or "")
    _autosize(ws)

    filename = (
        f"coach-{_slug(detail['coach_name'])}_"
        f"{window['date_from']}_to_{window['date_to']}.xlsx"
    )
    return {"filename": filename, "mime": XLSX_MIME, "content_base64": _to_b64(wb)}, None


# ── all-coaches workbook ──────────────────────────────────────────────────────

def build_all_coaches_workbook(date_from=None, date_to=None):
    overview = list_coach_analytics(date_from=date_from, date_to=date_to)
    window = overview["window"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Coaches summary"
    ws["A1"] = "Coach analytics — all coaches"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Period: {window['date_from']} to {window['date_to']}"

    headers = [
        "Coach", "Juniors", "Sessions run", "Billable 1-on-1", "1-on-1 attendees",
        "Attendance rate", "Avg level", "Avg handicap change", "Evaluations",
        "Below", "Meeting", "Exceeding", "Promotion recs",
    ]
    _write_header(ws, 4, headers)
    for r, c in enumerate(overview["coaches"], start=5):
        ev = c["evaluations"]
        ws.cell(row=r, column=1, value=c["coach_name"])
        ws.cell(row=r, column=2, value=c["junior_count"])
        ws.cell(row=r, column=3, value=c["sessions_run"])
        ws.cell(row=r, column=4, value=c["billable_one_on_one"]["count"])
        ws.cell(row=r, column=5, value=c["billable_one_on_one"]["attendees"])
        ws.cell(row=r, column=6, value=_pct(c["attendance"]["rate"]))
        ws.cell(row=r, column=7, value=c["avg_current_level"] if c["avg_current_level"] is not None else "")
        ws.cell(row=r, column=8, value=_hcp(c["avg_handicap_change"]))
        ws.cell(row=r, column=9, value=ev["total"])
        ws.cell(row=r, column=10, value=ev["below_expectation"])
        ws.cell(row=r, column=11, value=ev["meeting_expectation"])
        ws.cell(row=r, column=12, value=ev["exceeding_expectation"])
        ws.cell(row=r, column=13, value=ev["move_next_level"])
    _autosize(ws)

    filename = f"coach-analytics_all_{window['date_from']}_to_{window['date_to']}.xlsx"
    return {"filename": filename, "mime": XLSX_MIME, "content_base64": _to_b64(wb)}
