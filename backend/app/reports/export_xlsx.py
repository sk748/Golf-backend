"""
Native .xlsx export for the reports-for-analysis module.

Follows coach_analytics/export.py: workbooks are built server-side with
openpyxl and returned base64-encoded inside the standard {data} envelope as
{filename, mime, content_base64}.

The section-assembly functions (junior_report_sections /
programme_report_sections) are shared with export_pdf.py so both formats
render exactly the same rows — only the final rendering differs.
"""

from openpyxl import Workbook
from openpyxl.styles import Font

from app.coach_analytics.export import (
    XLSX_MIME, _TITLE_FONT, _autosize, _pct, _slug, _to_b64, _write_header,
)


def _blank(value):
    return "" if value is None else value


def _period(window):
    if not window["date_from"] and not window["date_to"]:
        return "Period: all records"
    return f"Period: {window['date_from'] or 'start'} to {window['date_to'] or 'today'}"


def _window_tag(window):
    if not window["date_from"] and not window["date_to"]:
        return "all"
    return f"{window['date_from'] or 'start'}_to_{window['date_to'] or 'today'}"


# ── shared section assembly (xlsx + pdf render the same rows) ─────────────────

def junior_report_sections(report):
    """[(title, headers|None, rows)] — headers None means bold label/value rows."""
    j = report["junior"]
    att = report["attendance"]
    ev = report["evaluations"]
    mix = ev["assessment_mix"]
    prog = report["handicap_progress"]
    comps = report["competitions"]
    reqs = report["competition_requirements"]

    summary = [
        ("Name", j["full_name"]),
        ("Band", _blank(j["band"])),
        ("Current level", j["current_level"]),
        ("Attendance present", att["present"]),
        ("Attendance absent", att["absent"]),
        ("Attendance excused", att["excused"]),
        ("Attendance total", att["total"]),
        ("Attendance rate", _pct(att["rate"])),
        ("Evaluations — below expectation", mix["below_expectation"]),
        ("Evaluations — meeting expectation", mix["meeting_expectation"]),
        ("Evaluations — exceeding expectation", mix["exceeding_expectation"]),
        ("Latest recommendation", _blank(ev["latest_recommendation"])),
        ("Signed cards (verified rounds)", prog["signed_cards"]),
        ("Target signed cards", prog["target_signed_cards"]),
        ("Avg 9-hole score", _blank(prog["avg_9_hole"])),
        ("Avg 18-hole score", _blank(prog["avg_18_hole"])),
        ("Ready for handicap", "Yes" if prog["ready_for_handicap"] else "No"),
        ("Competitions played", comps["competitions_played"]),
        ("Best gross score", _blank(comps["best_gross_score"])),
        ("Competitive rounds this month", reqs["competitive_rounds"]["this_month"]),
    ]

    eval_rows = [
        [
            e["report_month"], e["current_level"], e["assessment"],
            e["recommendation"], _blank(e["avg_score_9"]),
            _blank(e["avg_score_18"]), _blank(e["best_gross_score"]),
        ]
        for e in ev["items"]
    ]

    history_rows = [
        [
            h["date_played"], h["gross_score"], _blank(h["score_differential"]),
            _blank(h["handicap_after"]),
            "Yes" if h["counts_toward_handicap"] else "",
            h["status"],
        ]
        for h in report["handicap_history"]
    ]

    comp_rows = [
        [
            c["date"], c["tournament_name"], "internal", c["format"],
            _blank(c["gross_score"]), _blank(c["net_score"]),
            _blank(c["stableford_points"]), _blank(c["position"]),
        ]
        for c in comps["internal"]
    ] + [
        [
            _blank(x.get("date")), _blank(x.get("event_name")), "external",
            _blank(x.get("event_type")), _blank(x.get("gross_score")), "", "",
            _blank(x.get("position")),
        ]
        for x in comps["external"]
    ]

    return [
        ("Summary", None, summary),
        ("Evaluations",
         ["Month", "Level", "Assessment", "Recommendation", "Avg 9-hole", "Avg 18-hole", "Best gross"],
         eval_rows),
        ("Handicap history",
         ["Date", "Gross", "Differential", "Index after", "Counts", "Status"],
         history_rows),
        ("Competitions",
         ["Date", "Event", "Source", "Format", "Gross", "Net", "Stableford", "Position"],
         comp_rows),
    ]


def programme_report_sections(report):
    att = report["attendance"]
    hc = report["handicap"]
    tp = report["tournament_participation"]

    overview = [
        ("Juniors in programme", report["juniors_total"]),
        ("Attendance present", att["present"]),
        ("Attendance absent", att["absent"]),
        ("Attendance excused", att["excused"]),
        ("Attendance total", att["total"]),
        ("Attendance rate", _pct(att["rate"])),
        ("Current average handicap index", _blank(hc["current_avg_index"])),
        ("Internal tournament entries", tp["internal_entries"]),
        ("Juniors in internal tournaments", tp["internal_juniors"]),
        ("Verified external results", tp["verified_external_results"]),
        ("Juniors with external results", tp["external_juniors"]),
    ]

    band_rows = [[b["band"], b["juniors"]] for b in report["band_distribution"]]

    mix = report["evaluation_assessments"]
    eval_rows = [
        ["Below expectation", mix["below_expectation"]],
        ["Meeting expectation", mix["meeting_expectation"]],
        ["Exceeding expectation", mix["exceeding_expectation"]],
    ]

    trend_rows = [[t["month"], t["avg_handicap_index"]] for t in hc["monthly_avg_trend"]]

    return [
        ("Overview", None, overview),
        ("Band distribution", ["Band", "Juniors"], band_rows),
        ("Evaluation assessments", ["Assessment", "Count"], eval_rows),
        ("Handicap trend", ["Month", "Average index"], trend_rows),
    ]


# ── workbook rendering ────────────────────────────────────────────────────────

def _sections_workbook(title, period, sections):
    wb = Workbook()
    ws = wb.active
    for i, (name, headers, rows) in enumerate(sections):
        if i > 0:
            ws = wb.create_sheet()
        ws.title = name[:31]
        row = 1
        if i == 0:
            ws["A1"] = title
            ws["A1"].font = _TITLE_FONT
            ws["A2"] = period
            row = 4
        if headers is None:
            for r, (label, value) in enumerate(rows, start=row):
                ws.cell(row=r, column=1, value=label).font = Font(bold=True)
                ws.cell(row=r, column=2, value=value)
        else:
            _write_header(ws, row, headers)
            for r, values in enumerate(rows, start=row + 1):
                for c, value in enumerate(values, start=1):
                    ws.cell(row=r, column=c, value=value)
        _autosize(ws)
    return wb


def build_junior_workbook(report):
    window = report["window"]
    name = report["junior"]["full_name"]
    wb = _sections_workbook(
        f"Junior report — {name}", _period(window), junior_report_sections(report)
    )
    filename = f"junior-report-{_slug(name)}_{_window_tag(window)}.xlsx"
    return {"filename": filename, "mime": XLSX_MIME, "content_base64": _to_b64(wb)}


def build_programme_workbook(report):
    window = report["window"]
    wb = _sections_workbook(
        "Junior programme report", _period(window), programme_report_sections(report)
    )
    filename = f"programme-report_{_window_tag(window)}.xlsx"
    return {"filename": filename, "mime": XLSX_MIME, "content_base64": _to_b64(wb)}
